"""
Ingest PJM regulation signal sources into the local parquet store.

Sources (see CLAUDE.md Step 0 findings):

1. sample-normalize-new-regulation-signal-for-one-day.xlsx — the ONLY public
   post-Oct-2025 consolidated signal data: two real days (2025-10-03,
   2025-11-03) at 2s resolution, normalized ±1. Tidy layout: one sheet per
   day, columns Time / RegA.

2. rto-regulation-signal-data.zip — pre-redesign archive, rolling 12 months
   (2024-03 .. 2025-02 in the last posted version, ~484 MB, link removed from
   pjm.com after the redesign). One xlsx per month ("MM YYYY.xlsx"), two
   sheets: "Dynamic" (RegD — the fast, battery-oriented signal) and
   "Traditional" (RegA). Matrix layout: column A = time of day in 2s steps,
   row 1 = date per column, body = normalized signal. We ingest the Dynamic
   sheet as `signal` since a battery in the legacy regime followed RegD.
"""
from __future__ import annotations

import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path

import pandas as pd

from . import store

SAMPLE_XLSX_URL = ("https://www.pjm.com/-/media/DotCom/markets-ops/ancillary/"
                   "sample-normalize-new-regulation-signal-for-one-day.xlsx")
SIGNAL_ZIP_URL = ("https://www.pjm.com/-/media/DotCom/markets-ops/ancillary/"
                  "rto-regulation-signal-data.zip")

# A real day has 43,200 2-second points; source sheets often carry a stray
# midnight row that spills into the next day. Skip such stub days.
_MIN_ROWS_PER_DAY = 100


def ingest_sample_xlsx(path: str | Path) -> list[str]:
    """Load the post-redesign sample workbook (one sheet per day) into the store."""
    xl = pd.ExcelFile(path)
    written = []
    for sheet in xl.sheet_names:
        df = xl.parse(sheet)
        df.columns = [str(c).strip().lower() for c in df.columns]
        tcol = next(c for c in df.columns if "time" in c)
        scol = next(c for c in df.columns if c != tcol)
        out = pd.DataFrame({"signal": df[scol].astype(float).to_numpy()},
                           index=pd.DatetimeIndex(pd.to_datetime(df[tcol]), name="ts"))
        out = out.dropna()
        for day, day_df in out.groupby(out.index.normalize()):
            if len(day_df) < _MIN_ROWS_PER_DAY:
                continue
            store.write_signal_day(day, day_df)
            written.append(str(day.date()))
    return written


def _parse_monthly_matrix(xlsx_path: Path, sheet: str = "Dynamic") -> pd.DataFrame:
    """
    Parse one monthly matrix sheet into a tidy frame (index ts, column signal).
    Requires openpyxl cached formula values (data_only=True): row 1 holds the
    date per column, column A the time of day in 2-second steps.
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"{xlsx_path.name}: no sheet {sheet!r} (has {wb.sheetnames})")
    ws = wb[sheet]

    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    day_by_col: dict[int, date] = {}
    for j, v in enumerate(header[1:], start=1):
        if isinstance(v, datetime):
            day_by_col[j] = v.date()
        elif isinstance(v, date):
            day_by_col[j] = v

    times: list[time] = []
    values: list[tuple] = []
    for row in rows:
        t = row[0]
        if isinstance(t, datetime):
            t = t.time()
        if not isinstance(t, time):
            continue
        times.append(t)
        values.append(row)

    frames = []
    for j, d in day_by_col.items():
        sig = pd.to_numeric(pd.Series([r[j] if j < len(r) else None for r in values]),
                            errors="coerce")
        idx = pd.DatetimeIndex(
            [datetime.combine(d, t) for t in times], name="ts")
        frames.append(pd.DataFrame({"signal": sig.to_numpy()}, index=idx).dropna())
    wb.close()
    if not frames:
        return pd.DataFrame(columns=["signal"])
    out = pd.concat(frames).sort_index()
    return out[~out.index.duplicated(keep="first")]


def ingest_signal_zip(path: str | Path, member_filter: str | None = None,
                      sheet: str = "Dynamic",
                      workdir: str | Path | None = None) -> list[str]:
    """
    Ingest monthly members of the pre-redesign signal zip.

    member_filter: regex on member names, e.g. r'(01|02) 2025' for Jan-Feb 2025.
    sheet: "Dynamic" (RegD, default — what a battery followed) or "Traditional".
    Returns the list of days written.
    """
    path = Path(path)
    workdir = Path(workdir) if workdir else path.parent
    written: list[str] = []
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
        if member_filter:
            names = [n for n in names if re.search(member_filter, n)]
        for name in sorted(names):
            target = workdir / name
            if not target.exists():
                z.extract(name, workdir)
            tidy = _parse_monthly_matrix(target, sheet=sheet)
            for day, day_df in tidy.groupby(tidy.index.normalize()):
                if len(day_df) < _MIN_ROWS_PER_DAY:
                    continue
                store.write_signal_day(day, day_df)
                written.append(str(day.date()))
    return sorted(set(written))
